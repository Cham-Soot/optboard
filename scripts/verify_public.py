"""Reject private data before committing or publishing generated market files."""
import json
from pathlib import Path
import re
import subprocess
from public_data import assert_public

ROOT = Path(__file__).resolve().parent.parent


def verify(root=ROOT):
    for directory in ("data", "archive"):
        for path in (root / directory).glob("[0-9][0-9][0-9][0-9].json"):
            assert_public(json.loads(path.read_text(encoding="utf-8")), str(path.relative_to(root)))
    index = root / "index.html"
    if index.exists():
        html = index.read_text(encoding="utf-8")
        if not re.search(r"const MYDATA\s*=\s*null\s*;", html):
            raise ValueError("공개 화면에 개인 데이터가 내장되어 있거나 MYDATA 경계가 없습니다")
    tracked = subprocess.check_output(["git", "-c", "core.quotepath=false", "ls-files"], cwd=root, text=True, encoding="utf-8").splitlines()
    forbidden = [name for name in tracked if name in (".env", ".token_cache.json", "data/marks.json", "내화면.html")
                 or name.startswith((".private/", "node_modules/", "_site/")) or name.endswith(".local.json")]
    if forbidden:
        raise ValueError("개인 파일이 Git 추적 목록에 있습니다: " + ", ".join(forbidden))
    for path in (root / "archive").glob("*.xlsx"):
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=False)
        try:
            for sheet in book:
                for row in sheet.iter_rows(min_row=2, min_col=10, max_col=11, values_only=True):
                    if any(value not in (None, "") for value in row):
                        raise ValueError("공개 엑셀 보관본에 개인 메모가 있습니다: " + path.name)
        finally:
            book.close()
    print("공개 파일 검사 통과")


if __name__ == "__main__":
    verify()
