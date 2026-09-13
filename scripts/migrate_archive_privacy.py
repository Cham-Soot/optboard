"""Back up private annotations from legacy public Excel archives, then rebuild quotes only."""
import datetime as dt
import json
from pathlib import Path
import shutil
from openpyxl import load_workbook
from archive import write_xlsx, fmt_strike, write_index
from public_data import assert_public, public_doc, atomic_json

ROOT = Path(__file__).resolve().parent.parent


def extract(path, expiry):
    marks, memos = {}, {}
    book = load_workbook(path, data_only=False)
    try:
        for sheet in book:
            strike = fmt_strike(sheet.title)
            for row in sheet.iter_rows(min_row=2):
                if len(row) < 9 or not row[4].value:
                    continue
                date = row[4].value
                date = date.date().isoformat() if isinstance(date, dt.datetime) else str(date)[:10]
                dt.date.fromisoformat(date)
                key = expiry + "|" + strike + "|" + date
                pair = [str(row[i].value) if i < len(row) and row[i].value is not None else "" for i in (9, 10)]
                if any(pair): memos[key] = pair
                shape = {}
                for cell, i in (("cH",1),("cL",2),("pH",6),("pL",7)):
                    border = row[i].border
                    if border.bottom.style:
                        shape[cell] = "box" if all(getattr(border, side).style for side in ("top","left","right")) else "ul"
                if shape: marks[key] = shape
    finally:
        book.close()
    return {"formatVersion":2, "marks":marks, "memos":memos}


def main():
    private = ROOT / ".private" / "legacy-archive"
    private.mkdir(parents=True, exist_ok=True)
    for path in (ROOT / "archive").glob("*.xlsx"):
        recovered = extract(path, path.stem)
        if not recovered["marks"] and not recovered["memos"]:
            continue
        backup = private / path.name
        if backup.exists() and backup.read_bytes() != path.read_bytes():
            raise ValueError("이름이 같은 다른 백업이 있어 중단합니다: " + str(backup))
        shutil.copyfile(path, backup)
        atomic_json(private / (path.stem + "-memos.json"), recovered)
        doc = json.loads(path.with_suffix(".json").read_text(encoding="utf-8"))
        assert_public(doc, path.name)
        if not write_xlsx(public_doc(doc), {}, {}, path):
            raise RuntimeError("공개 엑셀 생성 실패")
        print("%s: 메모 %d행, 마킹 %d행 개인 백업 후 공개 보관본 재생성" %
              (path.name, len(recovered["memos"]), len(recovered["marks"])))
    write_index()


if __name__ == "__main__":
    main()
